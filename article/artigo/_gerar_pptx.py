"""Gera o PPTX da banca a partir do roteiro alinhado."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
FIG = ROOT / "article" / "artigo" / "figuras"
OUT = ROOT / "article" / "artigo" / "Pipeline_Churn_Segmentacao_Telecom.pptx"
DOWNLOADS = Path(r"C:\Users\matheus\Downloads\Pipeline_de_Churn_e_Segmentação_em_Telecom.pptx")

GREEN = RGBColor(0x0B, 0x5C, 0x36)
RED = RGBColor(0xB0, 0x1C, 0x2E)
DARK = RGBColor(0x2B, 0x2B, 0x2B)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xF4, 0xF6, 0xF5)
LINE = RGBColor(0xD5, 0xDD, 0xD8)
N0 = RGBColor(0x2E, 0x86, 0xAB)
N1 = RGBColor(0xE8, 0x48, 0x55)
N2 = RGBColor(0xF6, 0xAE, 0x2D)
FONT = "Calibri"


def set_run(run, size=18, bold=False, color=DARK, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {})
        rpr.append(ea)
    ea.set("typeface", font)
    latin = rpr.find(qn("a:latin"))
    if latin is None:
        latin = rpr.makeelement(qn("a:latin"), {})
        rpr.append(latin)
    latin.set("typeface", font)


def box(slide, l, t, w, h, fill=None, line=None, radius=True):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill if fill else WHITE
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    return sh


def textbox(slide, l, t, w, h, lines, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = line
        set_run(run, size=size, bold=bold, color=color)
    return tb


def title(slide, s, size=28):
    textbox(slide, 0.55, 0.28, 12.2, 0.85, [s], size=size, bold=True, color=GREEN)


def bullets(slide, l, t, w, h, items, size=18):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = "•  " + item
        set_run(run, size=size, color=DARK)
    return tb


def card_text(slide, l, t, w, h, heading, body, fill=CARD, head_color=GREEN):
    box(slide, l, t, w, h, fill=fill, line=LINE)
    textbox(slide, l + 0.16, t + 0.12, w - 0.32, 0.4, [heading], size=16, bold=True, color=head_color)
    tb = slide.shapes.add_textbox(Inches(l + 0.16), Inches(t + 0.52), Inches(w - 0.32), Inches(h - 0.64))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = body
    set_run(run, size=15, color=DARK)


def add_picture(slide, path, l, t, w, h):
    slide.shapes.add_picture(str(path), Inches(l), Inches(t), Inches(w), Inches(h))


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 capa
    s = new_slide(prs)
    box(s, 0, 0, 0.18, 7.5, fill=GREEN, radius=False)
    textbox(s, 0.7, 1.7, 11.8, 2.6, [
        "Pipeline para Predição de Churn e",
        "Segmentação de Clientes em Risco",
        "no Setor de Telecomunicações",
    ], size=32, bold=True, color=GREEN)
    textbox(s, 0.7, 4.6, 11.5, 1.8, [
        "Matheus Eduardo Laurindo do Rego",
        "Especialização em Ciência de Dados — UTFPR",
        "matheusrego@alunos.utfpr.edu.br",
        "Banca do Projeto Integrador",
    ], size=18, color=DARK)

    # 2 problema
    s = new_slide(prs)
    title(s, "No telecom, a receita recorrente foge sem atrito — e a lista de risco chega crua")
    box(s, 0.5, 1.35, 4.3, 5.5, fill=CARD, line=LINE)
    box(s, 0.85, 1.7, 1.35, 3.4, fill=RGBColor(0xB0, 0xB6, 0xB3), radius=False)
    textbox(s, 0.75, 5.2, 1.6, 0.9, ["1 lista", "1 tratamento"], size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    for i, (c, lab) in enumerate(((N0, "N0"), (N1, "N1"), (N2, "N2"))):
        box(s, 2.7, 1.75 + i * 1.15, 1.75, 0.95, fill=c)
        textbox(s, 2.7, 1.95 + i * 1.15, 1.75, 0.55, [lab], size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    textbox(s, 2.55, 5.25, 2.1, 0.9, ["3 núcleos", "3 ações"], size=13, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    bullets(s, 5.1, 1.4, 7.7, 5.6, [
        "O projeto parte do cenário das operadoras de telecomunicações (banda larga): receita recorrente, canais digitais e planos sem fidelidade, em que trocar de empresa ficou barato para o cliente.",
        "Cancelar (churn) custa caro: adquirir um assinante novo pesa bem mais do que reter quem já está na base.",
        "A resposta usual da Ciência de Dados é um classificador — uma lista de quem deve cancelar, sem dizer que tipo de cliente é cada um.",
        "Sem perfil, a operação trata todo cliente em risco da mesma forma — quando cada tipo pode exigir uma ação diferente e adaptada.",
    ], size=16)

    # 3 proposta
    s = new_slide(prs)
    title(s, "O Projeto Integrador é um pipeline mensal de núcleos")
    cards = [
        ("Classificar", "O XGBoost, com limiar 0,795, marca quem está em risco."),
        ("Agrupar", "O K-Means roda só nesses sinalizados, não na base inteira."),
        ("Devolver", "Cada núcleo traz volume, fatura média e receita exposta."),
    ]
    for i, (h, b) in enumerate(cards):
        card_text(s, 0.5 + i * 4.2, 1.4, 3.95, 3.3, h, b)
        if i < 2:
            textbox(s, 4.2 + i * 4.2, 2.6, 0.5, 0.4, ["→"], size=24, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    textbox(s, 0.55, 5.0, 12.2, 1.8, [
        "A cada ciclo a carteira entra e saem grupos de clientes em risco, já descritos em volume, fatura e receita.",
        "Sobre os núcleos, o estudo traz uma análise exemplar de campanha; oferta e KPI ficam com a operação.",
    ], size=17, color=DARK)

    # 4 piloto
    s = new_slide(prs)
    title(s, "O piloto mostrou sinal, mas não chegou nas metas")
    pills = [("Amostra", "1.000"), ("AUC", "0,787"), ("Acurácia", "75,5%"), ("Precisão", "55,3%"), ("F1", "46,2%")]
    for i, (k, v) in enumerate(pills):
        box(s, 0.5 + i * 2.5, 1.3, 2.35, 1.15, fill=CARD, line=LINE)
        textbox(s, 0.5 + i * 2.5, 1.38, 2.35, 0.35, [k], size=13, color=GRAY, align=PP_ALIGN.CENTER)
        textbox(s, 0.5 + i * 2.5, 1.7, 2.35, 0.55, [v], size=22, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    bullets(s, 0.55, 2.7, 12.2, 4.3, [
        "Recorte reduzido: 1.000 clientes, só tempo de casa e tipo de contrato, regressão logística e limiar 0,5.",
        "No teste: AUC 0,787, acurácia 75,5%, precisão 55,3% e F1 46,2% — cerca de 2 vezes a taxa base, ainda abaixo das metas de 80%.",
        "198 sinalizados em 2 grupos grosseiros (DSL vs fibra) e R$ 11,6 mil expostos: o problema é tratável; o recorte mínimo não basta.",
    ], size=18)

    # 5 arquitetura
    s = new_slide(prs)
    title(s, "O mesmo código roda de novo a cada carteira")
    steps = ["Carteira\ndo mês", "extract\n(Kaggle)", "train →\nchampion", "score", "núcleos"]
    for i, lab in enumerate(steps):
        box(s, 0.35 + i * 2.6, 1.35, 2.35, 1.7, fill=CARD, line=GREEN)
        slide_text = s.shapes.add_textbox(Inches(0.4 + i * 2.6), Inches(1.5), Inches(2.25), Inches(1.45))
        tf = slide_text.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = lab
        set_run(run, size=14, bold=True, color=DARK)
        if i < 4:
            textbox(s, 2.55 + i * 2.6, 1.9, 0.4, 0.4, ["→"], size=18, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    bullets(s, 0.55, 3.35, 12.2, 3.7, [
        "Os notebooks definem a receita (Optuna, SHAP, K = 3); o pacote operacional repete o lote no mês seguinte.",
        "Sem champion no MLflow, o score não pontua — só entra modelo aprovado.",
        "Os números a seguir (1.556 clientes, R$ 123,4 mil) são a corrida de referência do laboratório; cada ciclo devolve o mesmo tipo de tabela.",
    ], size=17)

    # 6 dado
    s = new_slide(prs)
    title(s, "Contrato mensal concentra o cancelamento")
    add_picture(s, FIG / "eda_contrato.png", 0.4, 1.25, 7.1, 5.4)
    textbox(s, 0.45, 6.7, 7.0, 0.5, ["Taxa de churn por contrato. Linha tracejada = média da base (26,5%)."], size=12, color=GRAY)
    bullets(s, 7.7, 1.35, 5.2, 5.6, [
        "A base pública IBM Telco (7.043 clientes) entra como proxy da operadora: a base real não pode ser divulgada.",
        "No contrato mensal o churn é 42,7%; no de dois anos, 2,8%. Quem cancela também tem menos tempo de casa (18 vs 38 meses).",
        "Só essas duas variáveis não bastaram no piloto; o estudo completo derivou 47 atributos a partir das 21 colunas originais.",
    ], size=16)

    # 7 método
    s = new_slide(prs)
    title(s, "O limiar pesa mais do que a escolha do algoritmo")
    box(s, 0.5, 1.3, 12.3, 1.25, fill=CARD, line=GREEN)
    textbox(s, 0.5, 1.55, 12.3, 0.8, ["S  =  0,6 · F(β = 0,5)  +  0,4 · acurácia"], size=26, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    bodies = [
        ("Comparação", "Foram comparados Random Forest, XGBoost e LightGBM (Optuna, 40 tentativas, 5 folds), com peso de classe e sem SMOTE."),
        ("Por que a precisão", "Cada nome sinalizado vira contato. Por isso o objetivo privilegia precisão e o limiar entra dentro da busca."),
        ("Quem vence", "Os três empatam; o XGBoost vence no escore composto. Limiar calibrado: 0,795 (medir em 0,5 seria enganoso)."),
    ]
    for i, (h, b) in enumerate(bodies):
        card_text(s, 0.5 + i * 4.2, 2.8, 4.0, 3.9, h, b)

    # 8 teste
    s = new_slide(prs)
    title(s, "Dois em cada três contatados cancelariam de fato")
    mets = [("AUC", "0,847"), ("Acurácia", "80,6%"), ("Precisão", "66,4%"), ("Recall", "54,0%"), ("F1", "59,6%")]
    for i, (k, v) in enumerate(mets):
        box(s, 0.4 + i * 2.55, 1.2, 2.4, 1.0, fill=CARD, line=LINE)
        textbox(s, 0.4 + i * 2.55, 1.25, 2.4, 0.3, [k], size=12, color=GRAY, align=PP_ALIGN.CENTER)
        textbox(s, 0.4 + i * 2.55, 1.52, 2.4, 0.5, [v], size=20, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    add_picture(s, FIG / "avaliacao_modelo.png", 0.4, 2.35, 8.3, 3.5)
    textbox(s, 0.45, 5.9, 8.2, 0.4, ["Teste com 1.409 clientes: matriz, ROC e teto de F1 × acurácia."], size=12, color=GRAY)
    bullets(s, 8.85, 2.35, 4.15, 4.6, [
        "De 374 cancelamentos reais: 202 certos, 102 alarmes falsos e 172 casos perdidos. Precisão ≈ 2,5 vezes a taxa base (26,5%).",
        "Meta de acurácia 80%: atingida. Meta de F1 80%: não — o teto desta base fica em torno de 64%.",
        "O limite está nos dados (sem consumo, atendimento ou qualidade), não na falta de um quarto algoritmo.",
    ], size=14)

    # 9 shap
    s = new_slide(prs)
    title(s, "O modelo se apoia nas features derivadas, não em demografia")
    add_picture(s, FIG / "shap_beeswarm.png", 0.35, 1.2, 7.0, 5.5)
    textbox(s, 0.4, 6.75, 7.0, 0.45, ["15 variáveis mais influentes. Azul = valor baixo; vermelho = alto."], size=12, color=GRAY)
    bullets(s, 7.55, 1.35, 5.4, 5.7, [
        "As duas variáveis mais importantes foram criadas no estudo: fatura × contrato mensal e o escore de risco. Em terceiro vem o tempo de casa.",
        "Tenure alto e contrato de dois anos reduzem o risco estimado; gênero não entra no topo.",
        "O SHAP explica o modelo, não a causa — não se afirma que baixar a fatura reduz o cancelamento.",
    ], size=16)

    # 10 núcleos
    s = new_slide(prs)
    title(s, "A saída de cada ciclo são três núcleos, não uma lista única")
    nuc = [
        (N0, "Núcleo 0 — novos em DSL", ["241 clientes (15,5%)", "2,7 meses · R$ 42,47/mês", "65,6% cancelaram"]),
        (N1, "Núcleo 1 — veteranos em fibra", ["427 clientes (27,4%)", "26,7 meses · R$ 94,54/mês", "66,5% cancelaram"]),
        (N2, "Núcleo 2 — novos em fibra", ["888 clientes (57,1%)", "4,8 meses · R$ 81,93/mês", "72,1% cancelaram"]),
    ]
    for i, (c, h, lines) in enumerate(nuc):
        box(s, 0.45 + i * 4.25, 1.35, 4.05, 4.5, fill=c)
        textbox(s, 0.6 + i * 4.25, 1.55, 3.75, 0.9, [h], size=16, bold=True, color=WHITE)
        for j, line in enumerate(lines):
            box(s, 0.7 + i * 4.25, 2.65 + j * 0.9, 3.55, 0.75, fill=WHITE)
            textbox(s, 0.7 + i * 4.25, 2.78 + j * 0.9, 3.55, 0.5, [line], size=16, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    textbox(s, 0.5, 6.05, 12.3, 1.1, [
        "1.556 sinalizados na corrida de referência (limiar 0,795, K = 3). O classificador montou a lista; o K-Means montou os grupos que a empresa lê.",
    ], size=16, color=DARK)

    # 11 receita
    s = new_slide(prs)
    title(s, "Volume de clientes não é volume de receita")
    add_picture(s, FIG / "receita_participacao.png", 0.35, 1.25, 6.5, 5.3)
    textbox(s, 0.4, 6.6, 6.5, 0.5, ["Fatia de cada núcleo nos R$ 123,4 mil mensais expostos."], size=12, color=GRAY)
    bullets(s, 7.1, 1.4, 5.8, 5.6, [
        "Os 1.556 sinalizados expõem R$ 123,4 mil por mês (cerca de R$ 1,48 milhão no ano se ninguém for retido).",
        "O núcleo 0 tem 15,5% dos nomes e só 8,3% do dinheiro; o núcleo 2 tem 57,1% dos nomes e 59,0% da receita.",
        "Uma campanha única para os 1.556 misturaria quem paga R$ 42 com quem paga R$ 95 — por isso os núcleos existem.",
    ], size=16)

    # 12 campanha
    s = new_slide(prs)
    title(s, "Leitura exemplar: um caminho possível por núcleo")
    camp = [
        (N0, "Núcleo 0 — DSL novo, fatura baixa", "Hipótese de atrito de adoção: integração por canal barato e convite ao contrato anual."),
        (N1, "Núcleo 1 — fibra veterana, sem suporte", "Conta alta: contato humano e um período de suporte incluso, porque o ticket paga o esforço."),
        (N2, "Núcleo 2 — fibra recente, 59% da receita", "Maior risco: prioridade em qualidade de sinal e oferta de permanência."),
    ]
    for i, (c, h, b) in enumerate(camp):
        box(s, 0.5, 1.3 + i * 1.45, 12.3, 1.3, fill=CARD, line=LINE)
        box(s, 0.5, 1.3 + i * 1.45, 0.18, 1.3, fill=c, radius=False)
        textbox(s, 0.9, 1.38 + i * 1.45, 11.6, 0.4, [h], size=16, bold=True, color=GREEN)
        textbox(s, 0.9, 1.8 + i * 1.45, 11.6, 0.65, [b], size=16, color=DARK)
    textbox(s, 0.5, 5.75, 12.3, 1.3, [
        "Análise exemplar. O Projeto Integrador é o pipeline que republica os núcleos; canal, oferta e KPI ficam com a operação.",
    ], size=16, color=GRAY)

    # 13 Q&A
    s = new_slide(prs)
    title(s, "Cinco perguntas que resumem o projeto")
    qa = [
        ("O que é o Projeto Integrador?", "Um pipeline mensal: classifica o risco de churn e devolve núcleos de quem foi sinalizado."),
        ("Qual o problema no cenário de telecom?", "A lista de quem vai cancelar chega sem perfil, e a operação trata todo cliente em risco do mesmo jeito — quando cada tipo pede uma ação diferente."),
        ("O que o piloto mostrou?", "Que o problema é tratável — e que tenure + contrato, sozinhos, não chegam nas metas nem em núcleos úteis."),
        ("Qual o resultado da corrida de referência?", "XGBoost com AUC 0,847 e precisão 66,4%; 1.556 sinalizados, K = 3 e R$ 123,4 mil/mês expostos."),
        ("A campanha é a entrega?", "Não: é análise exemplar. A entrega é o pipeline que republica os núcleos a cada ciclo."),
    ]
    for i, (q, a) in enumerate(qa):
        top = 1.2 + i * 1.15
        box(s, 0.5, top, 12.3, 1.05, fill=CARD, line=LINE)
        textbox(s, 0.7, top + 0.06, 12.0, 0.35, [q], size=15, bold=True, color=GREEN)
        textbox(s, 0.7, top + 0.42, 12.0, 0.55, [a], size=14, color=DARK)

    # 14 fecho
    s = new_slide(prs)
    title(s, "Aprendizados da pós-graduação neste projeto")
    bullets(s, 0.55, 1.3, 12.2, 3.6, [
        "Da formação em Ciência de Dados veio a cadeia completa: extração e qualidade da base, EDA com teste estatístico, engenharia de atributos, classificação, agrupamento e interpretabilidade (SHAP).",
        "A disciplina de projeto exigiu ir além do notebook: o lote extract → train → score, o manifesto Production e a escrita do piloto e do artigo em formato SBC.",
        "O aprendizado que pesa na prática: escolher métrica pelo custo do contato, admitir o teto da base (~64% de F1) e traduzir modelo em núcleos que a operação consegue ler.",
    ], size=17)
    textbox(s, 0.55, 5.0, 12.2, 0.7, [
        "Limitações: sem linha do tempo; silhueta 0,283 (corte útil, não grupos naturais); receita = exposição bruta, não perda esperada.",
    ], size=14, color=GRAY)
    textbox(s, 0.55, 5.75, 12.2, 1.4, [
        "Obrigado.",
        "Matheus Eduardo Laurindo do Rego  ·  Especialização em Ciência de Dados — UTFPR",
        "matheusrego@alunos.utfpr.edu.br",
    ], size=18, color=DARK)

    prs.save(OUT)
    prs.save(DOWNLOADS)
    print(OUT)
    print(DOWNLOADS)


if __name__ == "__main__":
    build()
